"""不变量用例 TC-24 ~ TC-27（docs/02-报表用例.md §4）。

这些用例守护的是"工程可信度"而非单个功能点：
原始文件只读、输出幂等、清洗守恒、三方对账。
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from conftest import NORMAL_RECORDS, make_raw, sha256_of, write_xlsx
from excel_report import aggregate, cleaning, loader, pipeline
from excel_report.cleaning import SALES_AMOUNT_COLUMN

ALL = ["日期", "销售员", "产品", "数量", "单价", "地区"]

DIRTY_RECORDS = [
    *NORMAL_RECORDS,
    dict(NORMAL_RECORDS[0]),  # 重复
    {**NORMAL_RECORDS[1], "销售员": None},  # 关键字段缺失
    {**NORMAL_RECORDS[2], "地区": "   "},  # 空白串缺失
    {**NORMAL_RECORDS[0], "数量": 0},  # 数量<=0
    {**NORMAL_RECORDS[1], "数量": -7},  # 数量<0
    {**NORMAL_RECORDS[2], "数量": "待定"},  # 非法数值
    {**NORMAL_RECORDS[0], "日期": "2024-13-45"},  # 非法日期
    {**NORMAL_RECORDS[1], "产品": None},  # 产品缺失（应保留）
    {**NORMAL_RECORDS[2], "日期": 45292},  # Excel 序列日期（应保留）
]


# --- TC-24 原始文件只读（AC-06 / R-01） --------------------------------------
def test_raw_file_is_never_modified(tmp_path: Path) -> None:
    """运行前后源文件 SHA256 必须完全一致 —— 这是 R-01 的客观证据。"""
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)
    before = sha256_of(path)
    before_mtime = path.stat().st_mtime_ns

    pipeline.run(path, tmp_path / "output" / "report.xlsx")

    assert sha256_of(path) == before, "原始文件被修改了！"
    assert path.stat().st_mtime_ns == before_mtime, "原始文件的时间戳未应变化"


def test_cleaned_rows_never_exceed_raw_rows(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)

    result = pipeline.run(path, tmp_path / "report.xlsx")

    assert result.stats.cleaned_rows <= result.stats.raw_rows


def test_loader_module_has_no_write_api() -> None:
    """静态审查：loader 里不得出现任何写文件的 API（结构性防线的自动化检查）。"""
    import inspect

    source = inspect.getsource(loader)
    for forbidden in ("to_excel", "ExcelWriter", "open(", ".write", "mkdir", "shutil"):
        assert forbidden not in source, f"loader 中不应出现写文件相关调用: {forbidden}"


# --- TC-25 幂等性 / 可复现（AC-07 / C-08） -----------------------------------
def test_output_is_bytewise_idempotent(tmp_path: Path) -> None:
    """同一输入连续运行两次，输出必须**逐字节**一致。

    这条断言之所以能成立，是因为 report 模块显式固定了两处 openpyxl 的"当前时间"：
    ``docProps/core.xml`` 的 created/modified，以及每个 zip 成员的修改时间戳。
    若去掉该处理，本用例会失败（这正是 C-08 的客观证据）。
    """
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)
    first = tmp_path / "run1.xlsx"
    second = tmp_path / "run2.xlsx"

    pipeline.run(path, first)
    pipeline.run(path, second)

    assert sha256_of(first) == sha256_of(second), "两次运行的输出不一致，存在非确定性因素"


def test_output_has_no_wall_clock_timestamps(tmp_path: Path) -> None:
    """输出文件内不得残留"当前时间"痕迹（否则不可能逐字节复现）。"""
    import zipfile

    path = write_xlsx(tmp_path / "sales_raw.xlsx", NORMAL_RECORDS, columns=ALL)
    output = tmp_path / "report.xlsx"
    pipeline.run(path, output)

    with zipfile.ZipFile(output) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
        assert "2024-01-01T00:00:00Z" in core
        # 所有 zip 成员时间戳都被固定为 1980-01-01
        assert {info.date_time for info in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}


def test_normalized_output_is_still_a_valid_xlsx(tmp_path: Path) -> None:
    """规范化时间戳后，文件必须仍能被 openpyxl/pandas 正常打开（不能改坏文件）。"""
    from openpyxl import load_workbook

    from excel_report.report import SHEET_ORDER

    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)
    output = tmp_path / "report.xlsx"
    result = pipeline.run(path, output)

    workbook = load_workbook(output)
    assert workbook.sheetnames == list(SHEET_ORDER), "重写 zip 后工作表结构不能丢失"
    assert workbook["按销售员"].freeze_panes == "A2", "重写 zip 后样式不能丢失"
    readback = pd.read_excel(output, sheet_name="明细清洗后", engine="openpyxl")
    assert len(readback) == result.stats.cleaned_rows


def test_statistics_are_deterministic(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)

    first = pipeline.run(path, tmp_path / "a.xlsx")
    second = pipeline.run(path, tmp_path / "b.xlsx")

    pd.testing.assert_frame_equal(first.detail, second.detail)
    pd.testing.assert_frame_equal(first.by_person, second.by_person)
    pd.testing.assert_frame_equal(first.by_region, second.by_region)
    pd.testing.assert_frame_equal(first.overview, second.overview)


# --- TC-26 清洗守恒 ----------------------------------------------------------
def test_cleaning_conservation_on_dirty_dataset() -> None:
    raw = make_raw(DIRTY_RECORDS)
    _, stats = cleaning.clean(raw)

    total = (
        stats.cleaned_rows
        + stats.dropped_missing
        + stats.dropped_nonpositive
        + stats.dropped_duplicate
        + stats.dropped_invalid
    )
    assert total == stats.raw_rows == len(DIRTY_RECORDS)


@pytest.mark.parametrize("seed", [1, 7, 42, 99])
def test_cleaning_conservation_random_datasets(seed: int) -> None:
    """随机数据集上也必须恒等 —— 防止只在固定样例上成立的"假正确"。"""
    import random

    rng = random.Random(seed)
    records = []
    for _ in range(60):
        record = dict(rng.choice(NORMAL_RECORDS))
        if rng.random() < 0.15:
            record["销售员"] = None
        elif rng.random() < 0.15:
            record["数量"] = rng.choice([0, -1, -20])
        elif rng.random() < 0.15:
            record["单价"] = rng.choice(["待定", "面议"])
        elif rng.random() < 0.15:
            record["日期"] = rng.choice(["2024-13-45", "not-a-date"])
        records.append(record)

    raw = make_raw(records)
    _, stats = cleaning.clean(raw)

    assert stats.raw_rows == stats.cleaned_rows + stats.dropped_total


# --- TC-27 三方对账 ----------------------------------------------------------
def test_three_way_reconciliation_on_dirty_dataset(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)

    result = pipeline.run(path, tmp_path / "report.xlsx")

    detail_sum = round(float(result.detail[SALES_AMOUNT_COLUMN].sum()), 2)
    person_sum = round(float(result.by_person[SALES_AMOUNT_COLUMN].sum()), 2)
    region_sum = round(float(result.by_region[SALES_AMOUNT_COLUMN].sum()), 2)
    monthly_sum = round(float(result.monthly[SALES_AMOUNT_COLUMN].sum()), 2)

    assert detail_sum == person_sum == region_sum == monthly_sum == result.total_sales


def test_overview_total_matches_detail(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)
    result = pipeline.run(path, tmp_path / "report.xlsx")

    as_dict = {
        str(row["指标"]): row["数值"]
        for _, row in result.overview.iterrows()
        if pd.notna(row["指标"]) and str(row["指标"]).strip()
    }

    assert round(float(as_dict["总销售额"]), 2) == result.total_sales


def test_excel_roundtrip_preserves_values(tmp_path: Path) -> None:
    """写出的报表被读回后，数值必须与内存中的一致（防止 openpyxl 写出丢精度）。"""
    path = write_xlsx(tmp_path / "sales_raw.xlsx", DIRTY_RECORDS, columns=ALL)
    output = tmp_path / "report.xlsx"

    result = pipeline.run(path, output)
    readback = pd.read_excel(output, sheet_name="明细清洗后", engine="openpyxl")

    assert len(readback) == len(result.detail)
    assert readback[SALES_AMOUNT_COLUMN].round(2).tolist() == (
        result.detail[SALES_AMOUNT_COLUMN].round(2).tolist()
    )


def test_aggregate_pure_function_does_not_mutate_input() -> None:
    """aggregate 是纯函数：不得原地修改传入的 DataFrame。"""
    detail, _ = cleaning.clean(make_raw(DIRTY_RECORDS))
    snapshot = detail.copy(deep=True)

    aggregate.by_salesperson(detail)
    aggregate.by_region(detail)
    aggregate.monthly_totals(detail)
    aggregate.overview(detail, cleaning.CleaningStats(raw_rows=1, cleaned_rows=1))

    pd.testing.assert_frame_equal(detail, snapshot)


def test_clean_does_not_mutate_input() -> None:
    raw = make_raw(DIRTY_RECORDS)
    snapshot = raw.copy(deep=True)

    cleaning.clean(raw)

    pd.testing.assert_frame_equal(raw, snapshot)
